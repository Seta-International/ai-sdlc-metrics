"""Re-create the workbook's charts with English titles. Some openpyxl versions
drop chart objects on save while others preserve them across load/save; to get a
deterministic result either way, add_charts() first clears each target sheet's
charts, then adds exactly the intended ones. The template builder AND every
export fill call this so the output always has these 3 charts and no duplicates."""
from openpyxl.chart import BarChart, LineChart, RadarChart, Reference


def add_charts(wb) -> None:
    proj = wb["8. Dashboard-Project"]
    port = wb["9. Dashboard-Portfolio"]
    proj._charts = []
    port._charts = []
    radar = RadarChart()
    radar.title = "Maturity — 5 dimensions"
    radar.add_data(Reference(proj, min_col=2, min_row=7, max_row=11), titles_from_data=False)
    radar.set_categories(Reference(proj, min_col=1, min_row=7, max_row=11))
    proj.add_chart(radar, "D6")

    line = LineChart()
    line.title = "Trend (usage / AI-PR % / CFR / rework)"
    line.add_data(Reference(proj, min_col=2, max_col=5, min_row=20, max_row=32),
                  titles_from_data=True)
    line.set_categories(Reference(proj, min_col=1, min_row=21, max_row=32))
    proj.add_chart(line, "D20")

    bar = BarChart()
    bar.title = "Overall level by project"
    bar.add_data(Reference(port, min_col=7, min_row=6, max_row=17), titles_from_data=False)
    bar.set_categories(Reference(port, min_col=1, min_row=6, max_row=17))
    port.add_chart(bar, "I5")
